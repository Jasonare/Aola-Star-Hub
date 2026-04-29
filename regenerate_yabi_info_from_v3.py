#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook


INPUT_XLSX = "yabi_100bt_processed_v3.xlsx"
INPUT_JSON = "亚比大全.json"
OUT_JSON = "yabi_info_regenerated_from_v3.json"
OUT_XLSX = "yabi_info_regenerated_from_v3.xlsx"


def norm_text(v: Any) -> str:
    if v is None:
        return ""
    return str(v).replace("\xa0", " ").strip()


def to_int(v: Any, default: Optional[int] = None) -> Optional[int]:
    s = norm_text(v)
    if not s:
        return default
    m = re.search(r"-?\d+", s)
    if not m:
        return default
    return int(m.group(0))


def parse_id_set(v: Any) -> List[int]:
    s = norm_text(v)
    if not s:
        return []
    return [int(x) for x in re.findall(r"\d+", s)]


def load_name_map(json_path: Path) -> Tuple[Dict[str, int], Dict[int, str]]:
    raw = json_path.read_bytes()
    obj = None
    for enc in ("utf-8", "utf-8-sig", "gb18030", "gbk"):
        try:
            obj = json.loads(raw.decode(enc))
            break
        except Exception:
            pass
    if obj is None:
        raise ValueError("无法解析亚比大全.json")

    name_to_id: Dict[str, int] = {}
    id_to_name: Dict[int, str] = {}

    def add(name: Any, pid: Any) -> None:
        n = norm_text(name)
        if not n:
            return
        i = to_int(pid, None)
        if i is None:
            return
        name_to_id[n] = i
        # 优先保留更短的主名
        if i not in id_to_name or len(n) < len(id_to_name[i]):
            id_to_name[i] = n

    if isinstance(obj, list):
        for it in obj:
            if not isinstance(it, dict):
                continue
            add(
                it.get("name") or it.get("名称") or it.get("亚比名") or it.get("title"),
                it.get("id") or it.get("ID") or it.get("petId") or it.get("pet_id") or it.get("编号"),
            )
    elif isinstance(obj, dict):
        for k, v in obj.items():
            if isinstance(v, (str, int, float)):
                add(k, v)
        for key in ("data", "list", "items", "亚比列表"):
            arr = obj.get(key)
            if isinstance(arr, list):
                for it in arr:
                    if not isinstance(it, dict):
                        continue
                    add(
                        it.get("name") or it.get("名称") or it.get("亚比名") or it.get("title"),
                        it.get("id") or it.get("ID") or it.get("petId") or it.get("pet_id") or it.get("编号"),
                    )

    return name_to_id, id_to_name


def find_col(headers: List[str], preferred: List[str], default_idx_1based: Optional[int] = None) -> Optional[int]:
    # exact
    for p in preferred:
        for i, h in enumerate(headers, start=1):
            if h == p:
                return i
    # contains
    for p in preferred:
        for i, h in enumerate(headers, start=1):
            if p in h:
                return i
    return default_idx_1based


def calc_ability_level_100(race: Dict[str, int], talent: int = 0, study: int = 0) -> Dict[str, int]:
    # 奥拉星常用能力计算（中性性格，不加成）
    level = 100
    hp = int(((race["hp"] * 2 + talent + study // 4) * level) // 100 + level + 10)
    atk = int(((race["atk"] * 2 + talent + study // 4) * level) // 100 + 5)
    de = int(((race["def"] * 2 + talent + study // 4) * level) // 100 + 5)
    sa = int(((race["spAtk"] * 2 + talent + study // 4) * level) // 100 + 5)
    sd = int(((race["spDef"] * 2 + talent + study // 4) * level) // 100 + 5)
    sp = int(((race["speed"] * 2 + talent + study // 4) * level) // 100 + 5)
    return {"hp": hp, "atk": atk, "def": de, "spAtk": sa, "spDef": sd, "speed": sp}


def main() -> None:
    root = Path(".")
    xlsx_path = root / INPUT_XLSX
    if not xlsx_path.exists():
        raise FileNotFoundError(f"未找到输入文件: {xlsx_path}")
    json_path = root / INPUT_JSON
    if not json_path.exists():
        # 兜底：取最大的 json
        cands = sorted(root.glob("*.json"), key=lambda p: p.stat().st_size, reverse=True)
        if not cands:
            raise FileNotFoundError("未找到亚比大全.json")
        json_path = cands[0]

    _, id_to_name = load_name_map(json_path)

    wb = load_workbook(xlsx_path, data_only=True)
    if len(wb.worksheets) < 3:
        raise ValueError("工作表不足3个，无法解析")
    ws_skill = wb.worksheets[1]  # 第二页
    ws_race = wb.worksheets[2]   # 第三页

    skill_headers = [norm_text(ws_skill.cell(1, c).value) for c in range(1, ws_skill.max_column + 1)]
    race_headers = [norm_text(ws_race.cell(1, c).value) for c in range(1, ws_race.max_column + 1)]

    # 技能表列定位
    c_skill_ids = 5
    c_skill_name = 8
    c_level = 9
    c_power = 10
    c_pp = 11
    c_acc = 12
    c_attr_type = 13
    c_desc = 14
    c_type = 16
    c_attr = 17

    # 种族值表列定位
    c_race_ids = find_col(race_headers, ["亚比ID"], 1)
    # 用户要求最后一行的 L~Q 对应 体力~速度
    c_hp, c_atk, c_def, c_spa, c_spd, c_spe = 12, 13, 14, 15, 16, 17

    skills_by_id: Dict[int, List[Dict[str, Any]]] = defaultdict(list)
    races_by_id: Dict[int, Dict[str, int]] = {}

    # 1) 解析技能（按亚比ID集合展开到每个ID）
    for r in range(2, ws_skill.max_row + 1):
        ids = parse_id_set(ws_skill.cell(r, c_skill_ids).value if c_skill_ids else "")
        if not ids:
            continue
        name = norm_text(ws_skill.cell(r, c_skill_name).value if c_skill_name else "")
        if not name:
            continue

        lv = to_int(ws_skill.cell(r, c_level).value if c_level else None, 0) or 0
        power = to_int(ws_skill.cell(r, c_power).value if c_power else None, 0) or 0
        pp = to_int(ws_skill.cell(r, c_pp).value if c_pp else None, 0) or 0
        acc = to_int(ws_skill.cell(r, c_acc).value if c_acc else None, 100)
        if acc is None:
            acc = 100

        attr = norm_text(ws_skill.cell(r, c_attr).value if c_attr else "")
        atk_type = norm_text(ws_skill.cell(r, c_type).value if c_type else "")
        if (not attr or not atk_type) and c_attr_type:
            combo = norm_text(ws_skill.cell(r, c_attr_type).value)
            if "/" in combo:
                p1, p2 = combo.split("/", 1)
                if not attr:
                    attr = norm_text(p1)
                if not atk_type:
                    atk_type = norm_text(p2)
        desc = norm_text(ws_skill.cell(r, c_desc).value if c_desc else "")

        rec = {
            "技能名称": name,
            "获取等级": lv,
            "威力": power,
            "PP值": pp,
            "命中率": acc,
            "技能属性": attr,
            "技能类型": atk_type,
            "技能描述": desc,
        }
        for i in ids:
            skills_by_id[i].append(rec)

    # 去重（保持顺序）
    for i, arr in list(skills_by_id.items()):
        seen = set()
        out = []
        for it in arr:
            key = (
                it["技能名称"], it["获取等级"], it["威力"], it["PP值"], it["命中率"],
                it["技能属性"], it["技能类型"], it["技能描述"],
            )
            if key in seen:
                continue
            seen.add(key)
            out.append(it)
        skills_by_id[i] = out

    # 2) 解析种族值：同一ID集合会重复，取出现的“最后一行”L~Q
    for r in range(2, ws_race.max_row + 1):
        ids = parse_id_set(ws_race.cell(r, c_race_ids).value if c_race_ids else "")
        if not ids:
            continue
        hp = to_int(ws_race.cell(r, c_hp).value, None)
        atk = to_int(ws_race.cell(r, c_atk).value, None)
        de = to_int(ws_race.cell(r, c_def).value, None)
        sa = to_int(ws_race.cell(r, c_spa).value, None)
        sd = to_int(ws_race.cell(r, c_spd).value, None)
        sp = to_int(ws_race.cell(r, c_spe).value, None)
        if None in (hp, atk, de, sa, sd, sp):
            continue
        race = {
            "hp": int(hp),
            "atk": int(atk),
            "def": int(de),
            "spAtk": int(sa),
            "spDef": int(sd),
            "speed": int(sp),
        }
        for i in ids:
            races_by_id[i] = race

    # 3) 汇总亚比信息 + 能力值
    all_ids = sorted(set(skills_by_id.keys()) | set(races_by_id.keys()))
    yabi_info: List[Dict[str, Any]] = []
    skill_rows: List[Dict[str, Any]] = []
    race_rows: List[Dict[str, Any]] = []
    ability_rows: List[Dict[str, Any]] = []

    for i in all_ids:
        race = races_by_id.get(i)
        ability = calc_ability_level_100(race, talent=0, study=0) if race else None
        item = {
            "亚比ID": i,
            "亚比名称": id_to_name.get(i, ""),
            "技能数量": len(skills_by_id.get(i, [])),
            "技能": skills_by_id.get(i, []),
            "种族值": race or {},
            "能力值(等级100_天赋0_学习力0)": ability or {},
        }
        yabi_info.append(item)

        for s in skills_by_id.get(i, []):
            row = {"亚比ID": i, "亚比名称": id_to_name.get(i, "")}
            row.update(s)
            skill_rows.append(row)

        if race:
            rr = {"亚比ID": i, "亚比名称": id_to_name.get(i, "")}
            rr.update(
                {
                    "体力": race["hp"],
                    "攻击": race["atk"],
                    "防御": race["def"],
                    "特攻": race["spAtk"],
                    "特防": race["spDef"],
                    "速度": race["speed"],
                }
            )
            race_rows.append(rr)
            ar = {"亚比ID": i, "亚比名称": id_to_name.get(i, "")}
            ar.update(
                {
                    "体力": ability["hp"],
                    "攻击": ability["atk"],
                    "防御": ability["def"],
                    "特攻": ability["spAtk"],
                    "特防": ability["spDef"],
                    "速度": ability["speed"],
                }
            )
            ability_rows.append(ar)

    # 输出 JSON
    (root / OUT_JSON).write_text(json.dumps(yabi_info, ensure_ascii=False, indent=2), encoding="utf-8")

    # 输出 Excel
    with pd.ExcelWriter(root / OUT_XLSX, engine="openpyxl") as writer:
        pd.DataFrame(yabi_info)[["亚比ID", "亚比名称", "技能数量"]].to_excel(writer, sheet_name="亚比总表", index=False)
        pd.DataFrame(skill_rows).to_excel(writer, sheet_name="技能表", index=False)
        pd.DataFrame(race_rows).to_excel(writer, sheet_name="种族值表", index=False)
        pd.DataFrame(ability_rows).to_excel(writer, sheet_name="能力值表", index=False)

    print("done")
    print(f"input_xlsx={xlsx_path.name}")
    print(f"input_json={json_path.name}")
    print(f"output_json={OUT_JSON}")
    print(f"output_xlsx={OUT_XLSX}")
    print(f"total_ids={len(all_ids)}")
    print(f"ids_with_skills={sum(1 for x in all_ids if x in skills_by_id)}")
    print(f"ids_with_race={sum(1 for x in all_ids if x in races_by_id)}")


if __name__ == "__main__":
    main()

