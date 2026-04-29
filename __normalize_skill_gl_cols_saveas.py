from openpyxl import load_workbook
from pathlib import Path
import re

src = Path(r"C:\工作娱乐\Aola-Star-Hub\原系亚比_100bt_编号_技能_种族值_skill_retry_v3.xlsx")
out = Path(r"C:\工作娱乐\Aola-Star-Hub\原系亚比_100bt_编号_技能_种族值_skill_retry_v3_GL校正.xlsx")
wb = load_workbook(src)
ws = wb[wb.sheetnames[1]]

COL = {'G':7,'H':8,'I':9,'J':10,'K':11,'L':12,'M':13,'N':14,'O':15,'P':16,'Q':17,'R':18,'X':24,'Y':25,'Z':26,'AA':27}


def norm(v):
    if v is None:
        return ""
    s = str(v).replace("\ufeff", "").replace("\xa0", " ").strip()
    return "" if s.lower() == "nan" else s


def to_int_text(v):
    s = norm(v)
    m = re.search(r"-?\d+", s)
    return m.group(0) if m else ""


def pick_first_non_empty(*vals):
    for v in vals:
        s = norm(v)
        if s:
            return s
    return ""


changed_rows = 0
cell_updates = 0
for r in range(2, ws.max_row + 1):
    old = {k: ws.cell(r, idx).value for k, idx in COL.items()}

    g_name = pick_first_non_empty(old['G'], old['M'])
    h_lv = pick_first_non_empty(to_int_text(old['H']), to_int_text(old['Q']), to_int_text(old['AA']))
    i_pow = pick_first_non_empty(to_int_text(old['I']))
    j_pp = pick_first_non_empty(to_int_text(old['J']), to_int_text(old['P']), to_int_text(old['Z']))

    attr_from_k = ""
    type_from_k = ""
    k_raw = norm(old['K'])
    if "/" in k_raw:
        a, b = k_raw.split("/", 1)
        attr_from_k = norm(a)
        type_from_k = norm(b)

    attr = pick_first_non_empty(attr_from_k, old['O'], old['X'])
    atk_type = pick_first_non_empty(type_from_k, old['N'], old['Y'])

    if attr and atk_type:
        k_attr_type = f"{attr}/{atk_type}"
    else:
        k_attr_type = pick_first_non_empty(k_raw, attr, atk_type)

    l_desc = pick_first_non_empty(old['L'], old['R'])

    new_vals = {'G':g_name,'H':h_lv,'I':i_pow,'J':j_pp,'K':k_attr_type,'L':l_desc}

    row_changed = False
    for k in ('G','H','I','J','K','L'):
        new_v = new_vals[k]
        cur_v = norm(old[k])
        if norm(new_v) != cur_v:
            ws.cell(r, COL[k]).value = new_v
            row_changed = True
            cell_updates += 1

    if row_changed:
        changed_rows += 1

wb.save(out)
print(f"saved={out}")
print(f"changed_rows={changed_rows}")
print(f"cell_updates={cell_updates}")
print(f"max_row={ws.max_row}")
